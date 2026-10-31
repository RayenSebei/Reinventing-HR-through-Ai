"""CV ingestion: extract what a human can SEE, report what is hidden, minimise PII.

Hidden text (white or near-white, tiny font, outside the page, zero-width
characters, hidden Word runs, instructions in metadata) is removed from the text
the LLM receives and reported to the recruiter. The candidate is not rejected for
it: the recruiter decides.
"""
import hashlib
import io
import zipfile

from . import guard
from .textutil import ZERO_WIDTH, mask_pii

ACCEPTED = ("pdf", "docx", "txt")


class IngestionError(ValueError):
    pass


def sniff_type(data: bytes, filename: str) -> str:
    """Decide the type from the bytes, not from the extension."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        raise IngestionError("Unsupported archive (only .docx is accepted)")
    try:
        data.decode("utf-8")
        return "txt"
    except UnicodeDecodeError as exc:
        raise IngestionError("Unsupported file type. Use PDF, DOCX or UTF-8 text.") from exc


# ---------------------------------------------------------------- PDF

def _is_whiteish(color) -> bool:
    if color is None:
        return False
    if isinstance(color, (int, float)):
        color = (color,)
    try:
        c = [float(x) for x in color]
    except (TypeError, ValueError):
        return False
    if len(c) == 1:                       # gray
        return c[0] >= 0.92
    if len(c) == 3:                       # RGB
        return min(c) >= 0.92
    if len(c) == 4:                       # CMYK
        return max(c) <= 0.08
    return False


def _hidden_reason(ch: dict, page) -> str | None:
    if ch.get("size", 10) < 3:
        return "tiny_font"
    if _is_whiteish(ch.get("non_stroking_color")):
        return "white_text"
    if ch["x1"] < 0 or ch["x0"] > page.width or ch["bottom"] < 0 or ch["top"] > page.height:
        return "off_page"
    return None


def _extract_pdf(data: bytes):
    import pdfplumber

    visible_pages, hidden = [], []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > 10:
            raise IngestionError("CV longer than 10 pages")
        meta_text = " ".join(str(v) for v in (pdf.metadata or {}).values() if isinstance(v, (str, bytes)))
        for page in pdf.pages:
            groups: dict[str, list[str]] = {}
            for ch in page.chars:
                reason = _hidden_reason(ch, page)
                if reason:
                    groups.setdefault(reason, []).append(ch["text"])
            for reason, chars in groups.items():
                text = "".join(chars).strip()
                if text:
                    hidden.append({"technique": reason, "page": page.page_number, "text": text[:500]})
            vis = page.filter(lambda o: o.get("object_type") != "char" or _hidden_reason(o, page) is None)
            visible_pages.append(vis.extract_text() or "")
    if meta_text and guard.scan(meta_text)["injection"]:
        hidden.append({"technique": "metadata", "page": None, "text": meta_text[:500]})
    return "\n".join(visible_pages), hidden


# ---------------------------------------------------------------- DOCX

def _extract_docx(data: bytes):
    import docx
    from docx.shared import Pt

    document = docx.Document(io.BytesIO(data))
    visible, hidden = [], []

    def handle_paragraph(par):
        line = []
        for run in par.runs:
            font = run.font
            reason = None
            if font.hidden:
                reason = "hidden_run"
            elif font.size is not None and font.size < Pt(3):
                reason = "tiny_font"
            elif font.color is not None and font.color.type is not None and str(font.color.rgb or "").upper() in {"FFFFFF", "FEFEFE", "FDFDFD"}:
                reason = "white_text"
            if reason:
                if run.text.strip():
                    hidden.append({"technique": reason, "page": None, "text": run.text[:500]})
            else:
                line.append(run.text)
        visible.append("".join(line))

    for par in document.paragraphs:
        handle_paragraph(par)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    handle_paragraph(par)
    core = document.core_properties
    meta_text = " ".join(filter(None, [core.title, core.subject, core.comments, core.keywords]))
    if meta_text and guard.scan(meta_text)["injection"]:
        hidden.append({"technique": "metadata", "page": None, "text": meta_text[:500]})
    return "\n".join(visible), hidden


# ---------------------------------------------------------------- entry point

def ingest_cv(data: bytes, filename: str, candidate_names: list[str]) -> dict:
    kind = sniff_type(data, filename)
    if kind == "pdf":
        visible, hidden = _extract_pdf(data)
    elif kind == "docx":
        visible, hidden = _extract_docx(data)
    else:
        visible, hidden = data.decode("utf-8"), []

    zero_width = len(ZERO_WIDTH.findall(visible))
    if zero_width:
        hidden.append({"technique": "zero_width_chars", "page": None, "text": f"{zero_width} invisible characters"})
    visible = ZERO_WIDTH.sub("", visible)

    if len(visible.strip()) < 40:
        raise IngestionError("No readable text found (scanned image?). Ask for a text PDF or DOCX.")

    visible_scan = guard.scan(visible)
    minimised, pii_counts = mask_pii(visible, candidate_names)

    return {
        "file_type": kind,
        "sha256": hashlib.sha256(data).hexdigest(),
        "text_minimised": minimised[:20000],
        "findings": {
            "hidden_content": hidden,
            "hidden_count": len(hidden),
            "visible_instruction_like_text": visible_scan["matches"] if visible_scan["injection"] else [],
            "pii_masked": pii_counts,
        },
    }
